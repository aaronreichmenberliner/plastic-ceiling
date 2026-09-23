#!/usr/bin/env python3
"""
Shared infrastructure for the per-figure build scripts.

Holds everything that more than one figure needs: input paths and their CLI overrides, rendering constants, the workbook readers, the compositing helpers, and the manifest writer. Nothing here draws a panel. Each figureN.py owns its own type sizes, panel geometry, panel generators, and canvas assembly.

Input paths are module globals so that a figure script can override them through configure() and the readers below pick the change up.
"""

from __future__ import annotations

import argparse
import base64
import hashlib
import io
import re
import json
import math
import os
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parent
os.environ.setdefault("MPLCONFIGDIR", str(ROOT / ".mplconfig"))

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl
from PIL import Image

# Default input/output locations (all overridable on the command line).
DATA = ROOT / "data"
ASSETS = ROOT / "assets"
OUTPUTS = ROOT / "outputs"
NUMBERS = DATA / "numbers.xlsx"
WORKBOOK = DATA / "bioprocess_sizing_framework.xlsx"
COLORS = DATA / "colors.xlsx"
FIG0 = ASSETS / "fig0_study_design.svg"
FIG2 = ASSETS / "fig2a_BPSF_flowchart.svg"
FIG5 = ASSETS / "fig5_bioprocess_sizing_framework.svg"

# Rendering constants.
# Helvetica first; TeX Gyre Heros and Nimbus Sans are metric-compatible clones. Resolving to a
# face that is actually installed matters twice over: matplotlib lays text out with it, and the
# SVG compositor names it, so layout metrics and render metrics agree. The previous silent
# fallback to DejaVu meant they did not.
FONT_STACK = ["Helvetica", "TeX Gyre Heros", "Nimbus Sans", "Liberation Sans", "Arial", "DejaVu Sans"]


def _resolve_font() -> str:
    from matplotlib import font_manager

    available = {f.name for f in font_manager.fontManager.ttflist}
    return next((name for name in FONT_STACK if name in available), "DejaVu Sans")


FONT = _resolve_font()
FONT_CSS = ", ".join(FONT_STACK)   # for <text> elements the compositor writes directly
HATCH_LW = 0.5        # hatch stroke width (pt); thinner strokes, denser repeats
KG_PER_TONNE = 1000.0
T_C = "t$_\\mathrm{C}$"     # tonnes of carbon; C set as a subscript index of t
KG_C = "kg$_\\mathrm{C}$"                 # kilograms of carbon, same convention
KG_BIOMASS = "kg$_\\mathrm{biomass}$"     # kilograms of dry biomass
DPI = 600             # panel rasterization density
OUT_DPI = 300         # dpi tag written into the final PNG/PDF
COLOR_SHEET = "Figure Color Key"
BASE_SIZE = 9.0       # rcParams default; each figure overrides per artist


# --------------------------------------------------------------------------- #
# Configuration and input checking
# --------------------------------------------------------------------------- #

def configure(**overrides: Path | None) -> None:
    """Rebind input/output globals from parsed CLI arguments. Readers in this module resolve them at call time, so a later override is picked up."""
    for name, value in overrides.items():
        if value is not None:
            globals()[name] = value.expanduser().resolve()
    OUTPUTS.mkdir(parents=True, exist_ok=True)


def add_path_arguments(parser: argparse.ArgumentParser, *names: str) -> None:
    """Attach the --numbers/--workbook/... flags a given figure actually consumes."""
    help_text = {
        "numbers": "Figure 1 values workbook.",
        "workbook": "Bioprocess sizing workbook.",
        "colors": "Shared color-key workbook.",
        "fig0": "Authored Figure 0 SVG.",
        "fig2": "Authored Figure 2 flowchart SVG.",
        "fig5": "Authored Figure 5 SVG.",
    }
    for name in names:
        parser.add_argument(f"--{name}", type=Path, default=None, help=help_text[name])
    parser.add_argument("--output-dir", type=Path, default=None, help="Directory for final figures.")


def check_inputs(paths: list[Path], number: int) -> None:
    """Fail fast and specifically, naming the missing file rather than dying part-way through compositing."""
    missing = [rel(path) for path in paths if not path.exists()]
    if missing:
        lines = [f"  {name}" for name in missing]
        raise SystemExit(f"Figure {number} is missing required input file(s):\n" + "\n".join(lines))


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(ROOT))
    except ValueError:
        return str(path)


# --------------------------------------------------------------------------- #
# Shared helpers
# --------------------------------------------------------------------------- #

def apply_rcparams() -> None:
    plt.rcParams.update(
        {
            "font.family": FONT,
            "font.size": BASE_SIZE,
            # Global because per-artist hatch linewidth needs matplotlib >= 3.10 and requirements.txt pins >= 3.8. Below ~0.35 pt the hatch aliases to flat gray at the composited scale.
            "hatch.linewidth": HATCH_LW,
            # Subscripts (t_C, kg_C, kg_biomass) must not drop into a different face.
            "mathtext.default": "regular",
            "mathtext.fontset": "custom",
            "mathtext.rm": FONT,
            "mathtext.it": f"{FONT}:italic",
            "mathtext.bf": f"{FONT}:bold",
            "mathtext.cal": FONT,
            # Panels are emitted as SVG and composited as SVG, so text stays live text.
            "svg.fonttype": "none",
        }
    )

def read_color_key(path: Path) -> dict[str, str]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[COLOR_SHEET]
    colors: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        label, hex_color = row[2], row[3]
        if isinstance(label, str) and isinstance(hex_color, str) and hex_color.startswith("#"):
            colors[label.strip()] = hex_color.strip()
    return colors

def read_pattern_key(path: Path) -> dict[str, str]:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[COLOR_SHEET]
    patterns: dict[str, str] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        label, pattern = row[2], row[4]
        if isinstance(label, str) and isinstance(pattern, str):
            patterns[label.strip()] = pattern.strip().lower()
    return patterns

def text_color(hex_color: str) -> str:
    hex_color = hex_color.lstrip("#")
    r, g, b = [int(hex_color[i : i + 2], 16) / 255 for i in (0, 2, 4)]
    luminance = 0.2126 * r + 0.7152 * g + 0.0722 * b
    return "black" if luminance > 0.62 else "white"

def display_formula(label: str) -> str:
    """Formula subscripts as mathtext rather than Unicode: the Helvetica-metric faces have no SUBSCRIPT TWO/FOUR glyph, and mathtext sets them from the body font."""
    return label.replace("CO2", "CO$_2$").replace("CH4", "CH$_4$")

def parse_route_label(label: str) -> tuple[str, str]:
    """Split a 'Substrate | Organism' scenario label into its two parts."""
    substrate, organism = [part.strip() for part in label.split("|", 1)]
    return substrate, organism

def organism_hatch(organism: str) -> str:
    """Hatch pattern for an organism or process class.

    Line motifs mark living catalysts, round motifs cell-free ones. An overlay marks a compound mode:
    the mixotroph's crosshatch is the phototroph's backslash over the heterotroph's forward slash, and
    is the only overlay in use. Thin routes take the vertical rule, whose strokes cross a stacked
    segment's fixed width; a horizontal rule's cross its height and vanish below one hatch cell (~4 mm).
    """
    if "Photoautotroph" in organism:
        return "\\\\\\"
    if "Lithoautotroph" in organism:
        return "|||"
    if "Mixotroph" in organism:
        return "xxx"
    if "Methanotroph" in organism:
        return "---"
    if "Enzymatic" in organism:
        return "..."
    if "Biochemical" in organism:
        return "ooo"
    return "///"

def clean_common(ax: plt.Axes, small: float = 7.5) -> None:
    """Axis styling for the plotted Figures 3 and 4 panels. Figure 4 passes size 8 tick labels; Figure 3 takes the 7.5 default."""
    ax.grid(False)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_color("#9A9A9A")
    ax.spines["bottom"].set_color("#9A9A9A")
    ax.tick_params(labelsize=small, colors="black")
    ax.set_facecolor("none")

def panel_title(ax: plt.Axes, title: str, subtitle: str, size: float, align: str = "left") -> None:
    """Bold title with an italic subtitle beneath it, above the axes. `align` is "left" or "center"; a centered title is right for a pie, whose content is centered on the axes, and wrong for a plot whose content starts at the y-axis."""
    x = 0.0 if align == "left" else 0.5
    ha = "left" if align == "left" else "center"
    ax.set_title(title, fontfamily=FONT, fontsize=size + 1.5, fontweight="bold", loc=align, pad=2.4 * size)
    ax.text(x, 1.015, subtitle, transform=ax.transAxes, ha=ha, va="bottom", fontfamily=FONT, fontsize=size, fontstyle="italic", color="#48566A")

def render(fig: plt.Figure, pad: float) -> str:
    """Serialize a panel figure to SVG. Vector, not raster: the composited figure keeps curves as curves and text as text, so the PDF has no resolution ceiling."""
    buf = io.BytesIO()
    fig.savefig(buf, format="svg", transparent=True, facecolor="none", edgecolor="none", bbox_inches="tight", pad_inches=pad)
    plt.close(fig)
    return buf.getvalue().decode()

_SANS_ALIASES = ("helvetica", "arial", "liberation sans", "nimbus sans", "tex gyre heros", "helvetica neue", "sans-serif")
# The opening quote, if any, belongs to the prefix; the value stops at the next quote or semicolon, so the
# quote that closes the enclosing attribute is never consumed and never has to be put back.
_FONT_FAMILY_DECL = re.compile(r'(font-family\s*[:=]\s*"?)([^";]+)')


def _normalise_font_family(svg: str) -> str:
    """Point every sans-serif font-family declaration in authored artwork at the package stack.

    Authoring tools emit whichever stack the source document carried, and a stack headed by a family that is
    not installed resolves to whatever clone is, so two panels of one figure can end up in different
    letterforms. Declarations that name something other than a Helvetica or Arial clone are left alone.
    """
    def rewrite(match: re.Match[str]) -> str:
        prefix, value = match.groups()
        if not any(alias in value.lower() for alias in _SANS_ALIASES):
            return match.group(0)
        return f"{prefix}{FONT_CSS}"

    return _FONT_FAMILY_DECL.sub(rewrite, svg)


_FACE_CACHE: dict[str, tuple] = {}
_TEXT_ELEMENT = re.compile(r"<text\b.*?</text>", re.S)
_TSPAN_OPEN = re.compile(r"<tspan\b[^>]*>")
_ATTR_OR_STYLE = r'(?:{name}\s*=\s*"([^"]*)"|{name}\s*:\s*([^;"]+))'


def _face(bold: bool) -> tuple:
    """Metrics for the resolved body face, loaded once. The file is located through matplotlib, so it follows FONT and whatever is installed rather than a fixed system path."""
    key = "bold" if bold else "regular"
    if key not in _FACE_CACHE:
        from fontTools.ttLib import TTFont
        from matplotlib import font_manager

        path = font_manager.findfont(font_manager.FontProperties(family=FONT, weight="bold" if bold else "normal"))
        try:
            face = TTFont(path, lazy=True)
        except Exception:
            # macOS ships Helvetica as a .ttc collection; take the first face in it.
            face = TTFont(path, lazy=True, fontNumber=0)
        _FACE_CACHE[key] = (face.getBestCmap(), face["hmtx"], face["head"].unitsPerEm)
    return _FACE_CACHE[key]


def _lookup(fragment: str, name: str) -> str | None:
    match = re.search(_ATTR_OR_STYLE.format(name=name), fragment)
    if not match:
        return None
    return (match.group(1) or match.group(2) or "").strip()


def _advance(text: str, size: float, bold: bool) -> float:
    cmap, hmtx, upm = _face(bold)
    total = 0
    for char in text:
        glyph = cmap.get(ord(char))
        if glyph is not None:
            total += hmtx[glyph][0]
    return total * size / upm


def text_advance_mm(text: str, font_pt: float, bold: bool = True) -> float:
    """Printed width of a run set in the body face, in mm, from the font's own advance widths."""
    return _advance(text, font_pt, bold) * 25.4 / 72.0


def _normalise_text_anchor(svg: str) -> str:
    """Re-express centred and right-aligned text as left-aligned text at an equivalent x.

    CairoSVG loses part of the advance at a tspan boundary when the element is centred, so a run split into
    tspans, as any run carrying a subscript must be, comes out narrower than it should and the words either
    side of the boundary close up. Left-aligned text is laid out correctly, so the anchor is resolved here
    instead: the width is measured from the font's own advances and folded into x. Only elements that are
    actually split are touched, and the geometry is unchanged.
    """
    def rewrite(match: re.Match[str]) -> str:
        element = match.group(0)
        head = element[: element.index(">") + 1]
        anchor = _lookup(head, "text-anchor")
        if anchor not in ("middle", "end"):
            return element
        x_attr = re.search(r'\bx\s*=\s*"([^"]+)"', head)
        if not x_attr or len(x_attr.group(1).split()) != 1:
            return element
        base_size = float((_lookup(head, "font-size") or "0").replace("px", "") or 0)
        base_bold = (_lookup(head, "font-weight") or "") == "bold"
        if not base_size:
            return element
        # A tspan carrying its own absolute x or y starts a fresh text chunk, which is anchored on its own and
        # is laid out correctly already. Such an element is a stack of lines, not one run, so its width is not
        # the sum of its parts and it must be left alone.
        if any(re.search(r'\s(?:x|y)\s*=\s*"', tag) for tag in _TSPAN_OPEN.findall(element)):
            return element

        # Walk the element as alternating markup and character data, tracking the size and weight in force.
        runs: list[tuple[str, float, bool]] = []
        size, bold = base_size, base_bold
        stack: list[tuple[float, bool]] = []
        for part in re.split(r"(<[^>]*>)", element[element.index(">") + 1 : element.rindex("</text>")]):
            if part.startswith("<"):
                if _TSPAN_OPEN.match(part):
                    stack.append((size, bold))
                    raw = _lookup(part, "font-size")
                    if raw:
                        value = raw.replace("px", "").strip()
                        size = size * float(value[:-2]) if value.endswith("em") else float(value)
                    weight = _lookup(part, "font-weight")
                    if weight:
                        bold = weight == "bold"
                elif part.startswith("</tspan") and stack:
                    size, bold = stack.pop()
            elif part:
                runs.append((part, size, bold))
        if len(runs) < 2:
            return element

        width = sum(_advance(t, s, b) for t, s, b in runs)
        x = float(x_attr.group(1))
        shifted = x - (width / 2.0 if anchor == "middle" else width)
        head_out = re.sub(r'(\bx\s*=\s*")[^"]+(")', lambda m: f"{m.group(1)}{shifted:.4f}{m.group(2)}", head, count=1)
        head_out = re.sub(_ATTR_OR_STYLE.format(name="text-anchor"), lambda m: 'text-anchor="start"' if m.group(1) is not None else "text-anchor:start", head_out, count=1)
        return head_out + element[element.index(">") + 1 :]

    return _TEXT_ELEMENT.sub(rewrite, svg)
_SUBSCRIPT_DIGITS = {"\u2080": "0", "\u2081": "1", "\u2082": "2", "\u2083": "3", "\u2084": "4", "\u2085": "5", "\u2086": "6", "\u2087": "7", "\u2088": "8", "\u2089": "9"}
_SUBSCRIPT_RUN = re.compile("[" + "".join(_SUBSCRIPT_DIGITS) + "]+")
# Shift and size of a synthesized subscript, as a fraction of the surrounding text. Expressed in em so the
# rewrite needs no knowledge of the font-size in force, which may be inherited or overridden by a style
# attribute. dy on the shifted tspan resolves against its own reduced size, so 0.28 em there and -0.2 em on
# the full-size reset both come to 0.2 em of the surrounding text.
_SUB_SIZE, _SUB_DOWN, _SUB_UP = "0.714em", "0.28em", "-0.2em"


def _normalise_subscripts(svg: str) -> str:
    """Rewrite Unicode subscript digits as shifted plain digits inside <text> elements.

    The Helvetica-metric faces carry no SUBSCRIPT TWO or SUBSCRIPT FOUR glyph, so a renderer falls back to
    another family for those characters alone, which changes their size and advance and leaves the run
    visibly misspaced. Plain digits in a shifted tspan are set from the same face as their surroundings.
    """
    def rewrite(match: re.Match[str]) -> str:
        def shift(run: re.Match[str]) -> str:
            digits = "".join(_SUBSCRIPT_DIGITS[c] for c in run.group(0))
            return f'<tspan font-size="{_SUB_SIZE}" dy="{_SUB_DOWN}">{digits}</tspan><tspan dy="{_SUB_UP}"></tspan>'
        parts = re.split(r"(<[^>]*>)", match.group(0))
        return "".join(p if p.startswith("<") else _SUBSCRIPT_RUN.sub(shift, p) for p in parts)

    return re.sub(r"<text\b.*?</text>", rewrite, svg, flags=re.S)


def render_svg(path: Path) -> str:
    """Authored SVG artwork, passed to the compositor with fonts and subscripts normalised. Nothing is rasterized, so <marker> arrowheads and stroke-dasharray survive."""
    return _normalise_text_anchor(_normalise_subscripts(_normalise_font_family(path.read_text(encoding="utf-8"))))

def open_rgba(path: Path) -> Image.Image:
    return Image.open(path).convert("RGBA")

def figure_with_plot_area(
    plot_width: float,
    plot_height: float,
    *,
    left: float = 0.0,
    right: float = 0.0,
    bottom: float = 0.0,
    top: float = 0.0,
) -> tuple[plt.Figure, plt.Axes]:
    """Figure whose axes occupy an exact physical plot area (Figure 1 pies)."""
    fig_width = left + plot_width + right
    fig_height = bottom + plot_height + top
    fig, ax = plt.subplots(figsize=(fig_width, fig_height), dpi=DPI, facecolor="none")
    ax.set_position([left / fig_width, bottom / fig_height, plot_width / fig_width, plot_height / fig_height])
    ax.set_facecolor("none")
    return fig, ax

def _trim_ink(image: Image.Image, pad: int) -> Image.Image:
    """Crop to the ink bounding box plus pad, so a composited panel carries no dead margin and its type is not scaled down by empty space."""
    ink = np.array(image.convert("L")) < 245
    rows, cols = np.where(ink.any(axis=1))[0], np.where(ink.any(axis=0))[0]
    box = (max(int(cols.min()) - pad, 0), max(int(rows.min()) - pad, 0), min(int(cols.max()) + pad + 1, image.width), min(int(rows.max()) + pad + 1, image.height))
    return image.crop(box)


# --------------------------------------------------------------------------- #
# Compositing helpers
# --------------------------------------------------------------------------- #

@dataclass
class Canvas:
    """An SVG document under construction. Coordinates are canvas pixels at OUT_DPI."""

    width: int
    height: int
    parts: list[str] = field(default_factory=list)


def _split_svg(svg: str) -> tuple[str, str, str]:
    """Return (root attributes, viewBox, body) for an SVG document."""
    svg = re.sub(r"<\?xml[^>]*\?>", "", svg)
    svg = re.sub(r"<!DOCTYPE[^>]*>", "", svg, flags=re.S)
    match = re.search(r"<svg\b[^>]*>", svg, re.S)
    root = match.group(0)
    body = svg[match.end() : svg.rindex("</svg>")]
    view_box = re.search(r'viewBox\s*=\s*"([^"]+)"', root)
    if view_box:
        box = view_box.group(1)
    else:
        width = re.search(r'\bwidth\s*=\s*"([\d.]+)', root).group(1)
        height = re.search(r'\bheight\s*=\s*"([\d.]+)', root).group(1)
        box = f"0 0 {width} {height}"
    attrs = re.sub(r'\s(?:x|y|width|height|viewBox|preserveAspectRatio)\s*=\s*"[^"]*"', "", root[4:-1])
    return attrs, box, body


def svg_size(svg: str) -> tuple[float, float]:
    """Width and height of a rendered panel in points, as `render` wrote them."""
    width = re.search(r'width="([\d.]+)pt"', svg)
    height = re.search(r'height="([\d.]+)pt"', svg)
    if not (width and height):
        raise ValueError("panel SVG carries no pt width/height; it did not come from render()")
    return float(width.group(1)), float(height.group(1))


def place(canvas: Canvas, item: str | Image.Image, box: tuple[int, int, int, int]) -> None:
    """Fit an SVG panel or a raster image into `box`, preserving aspect and centering."""
    x, y, width, height = box
    if isinstance(item, Image.Image):
        scale = min(width / item.width, height / item.height)
        w, h = item.width * scale, item.height * scale
        buf = io.BytesIO()
        item.save(buf, format="PNG")
        data = base64.b64encode(buf.getvalue()).decode()
        canvas.parts.append(f'<image x="{x + (width - w) / 2:.3f}" y="{y + (height - h) / 2:.3f}" width="{w:.3f}" height="{h:.3f}" xlink:href="data:image/png;base64,{data}"/>')
        return
    attrs, view_box, body = _split_svg(item)
    canvas.parts.append(f'<svg {attrs} x="{x}" y="{y}" width="{width}" height="{height}" viewBox="{view_box}" preserveAspectRatio="xMidYMid meet">{body}</svg>')


def text(canvas: Canvas, string: str, center: tuple[float, float], size: float, weight: str = "normal", color: str = "black") -> None:
    """Draw a horizontally centred text label onto the canvas, in canvas pixels, using the shared font stack."""
    x, y = center
    escaped = string.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    canvas.parts.append(f'<text x="{x:.3f}" y="{y:.3f}" font-family="{FONT_CSS}" font-size="{size:.3f}" font-weight="{weight}" fill="{color}" text-anchor="middle">{escaped}</text>')


def save_figure(canvas: Canvas, number: int, width_mm: float | None = None) -> None:
    import cairosvg

    OUTPUTS.mkdir(parents=True, exist_ok=True)
    png_path = OUTPUTS / f"Fig{number}.png"
    pdf_path = OUTPUTS / f"Fig{number}.pdf"
    # Physical output size. width_mm forces the final reproduction width (e.g. 180 mm double column) so the
    # figure prints 1:1 with no journal reduction; otherwise the canvas_px / OUT_DPI size is used.
    phys_w_in = (width_mm / 25.4) if width_mm is not None else (canvas.width / OUT_DPI)
    phys_h_in = phys_w_in * canvas.height / canvas.width
    png_dpi = canvas.width / phys_w_in
    document = (
        f'<svg xmlns="http://www.w3.org/2000/svg" xmlns:xlink="http://www.w3.org/1999/xlink" '
        f'width="{phys_w_in}in" height="{phys_h_in}in" viewBox="0 0 {canvas.width} {canvas.height}">'
        f'<rect x="0" y="0" width="{canvas.width}" height="{canvas.height}" fill="white"/>'
        + "".join(canvas.parts)
        + "</svg>"
    ).encode()
    cairosvg.svg2pdf(bytestring=document, write_to=str(pdf_path))
    png = cairosvg.svg2png(bytestring=document, output_width=canvas.width, output_height=canvas.height, background_color="white")
    Image.open(io.BytesIO(png)).convert("RGB").save(png_path, dpi=(png_dpi, png_dpi), optimize=True)
    print(f"Wrote {rel(png_path)} and {rel(pdf_path)}")


def read_sankey_flows() -> list[tuple[str, str, float]]:
    """Figure 1 Sankey flow table (source, target, carbon flow in kg_C/yr), from the 'Figure 1A' sheet."""
    wb = openpyxl.load_workbook(NUMBERS, data_only=True, read_only=True)
    ws = wb["Figure 1A"]
    flows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        source, target, value = row[0], row[1], row[2]
        if isinstance(source, str) and isinstance(target, str) and isinstance(value, (int, float)):
            flows.append((source.strip(), target.strip(), float(value)))
    return flows

def read_uplink_masses(categories: list[str]) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Total, carbon and non-carbon uplinked mass per stream, from the 'Figure 1C' sheet."""
    wb = openpyxl.load_workbook(NUMBERS, data_only=True, read_only=True)
    ws = wb["Figure 1C"]
    rows = {str(row[0]).strip(): row for row in ws.iter_rows(min_row=2, values_only=True) if row[0]}
    total, carbon, noncarbon = [], [], []
    for category in categories:
        row = rows[category]
        total.append(row[1])
        carbon.append(row[2])
        noncarbon.append(row[3])
    return np.array(total, float), np.array(carbon, float), np.array(noncarbon, float)

def read_terminal_waste() -> tuple[list[str], np.ndarray]:
    """Lifetime terminal carbon per waste stream, from the 'Figure 1B' sheet."""
    wb = openpyxl.load_workbook(NUMBERS, data_only=True, read_only=True)
    ws = wb["Figure 1B"]
    labels, values = [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        label = row[0]
        if isinstance(label, str) and label.strip() == "Total":
            break  # stop at a summary row so it is not ingested as a stream (same convention as read_scenario_carbon)
        if label and isinstance(row[1], (int, float)):
            labels.append(str(label).strip())
            values.append(float(row[1]))
    return labels, np.array(values, float)

def read_iss_input_carbon() -> float:
    """Total ISS carbon input (Inputs sheet), used by Figure 2."""
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = wb["Inputs"]
    for row in ws.iter_rows(values_only=True):
        if row[0] is None and isinstance(row[2], (int, float)):
            return float(row[2])
    raise ValueError("Could not find total input carbon in Inputs sheet.")

def read_scenario_carbon() -> tuple[list[str], np.ndarray, np.ndarray]:
    """Baseline (col 1) and aspirational (col 6) routed carbon per route (Figure 2)."""
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = wb["Scenarios"]
    labels, conservative, aggressive = [], [], []
    for row in ws.iter_rows(min_row=5, values_only=True):
        label = row[0]
        if label == "Total":
            break
        if not label:
            continue
        cons = row[1] if isinstance(row[1], (int, float)) else 0.0
        aggr = row[6] if isinstance(row[6], (int, float)) else 0.0
        labels.append(str(label))
        conservative.append(float(cons))
        aggressive.append(float(aggr))
    return labels, np.array(conservative), np.array(aggressive)

def read_scenario_rows() -> list[dict[str, float | str]]:
    """Full per-route record: carbon, biomass, and volume for both scenarios (Figure 3)."""
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = wb["Scenarios"]

    def num(value: object) -> float:
        return float(value) if isinstance(value, (int, float)) and math.isfinite(float(value)) else 0.0

    rows: list[dict[str, float | str]] = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        label = row[0]
        if label == "Total":
            break
        if not label:
            continue
        substrate, mechanism = parse_route_label(str(label))
        rows.append(
            {
                "label": str(label),
                "substrate": substrate,
                "mechanism": mechanism,
                "cons_carbon": num(row[1]),
                "cons_biomass_t": num(row[2]),
                "cons_biomass_e": num(row[3]),
                "cons_volume_t": num(row[4]),
                "cons_volume_e": num(row[5]),
                "aggr_carbon": num(row[6]),
                "aggr_biomass_t": num(row[7]),
                "aggr_biomass_e": num(row[8]),
                "aggr_volume_t": num(row[9]),
                "aggr_volume_e": num(row[10]),
            }
        )
    return rows

def read_fate_total() -> float:
    """Total routed carbon input for Figure 4 (Scenarios col 15, rows 6-9)."""
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = wb["Scenarios"]
    values = [ws.cell(row, 15).value for row in range(6, 10)]
    return float(sum(value for value in values if isinstance(value, (int, float))))

def read_fate_rows() -> dict[str, dict[str, float | str]]:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = wb["Scenarios"]
    rows: dict[str, dict[str, float | str]] = {}
    for row_idx in range(6, 10):
        label = ws.cell(row_idx, 14).value
        if not isinstance(label, str):
            continue
        rows[label] = {
            "baseline_theoretical": float(ws.cell(row_idx, 15).value),
            "aspirational_theoretical": float(ws.cell(row_idx, 17).value),
            "baseline_empirical": float(ws.cell(row_idx, 19).value),
            "aspirational_empirical": float(ws.cell(row_idx, 21).value),
            "fate": str(ws.cell(row_idx, 23).value),
        }
    return rows

def read_routing_rows() -> dict[str, tuple[float, float]]:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=True, read_only=True)
    ws = wb["Routing"]
    rows: dict[str, tuple[float, float]] = {}
    for row in ws.iter_rows(min_row=19, max_row=39, values_only=True):
        label, baseline, aspirational = row[0], row[1], row[2]
        if isinstance(label, str) and isinstance(baseline, (int, float)) and isinstance(aspirational, (int, float)):
            rows[label] = (float(baseline), float(aspirational))
    return rows


# --------------------------------------------------------------------------- #
# Legend tables shared by Figures 2 and 3
# --------------------------------------------------------------------------- #

# ---------------------------------------------------------------------------- study-wide route palette
# Approved similarity-ordered, colourblind-safe palette (gas -> soluble -> polymer families). CO2 carries
# two shades for its two routes (Photoautotroph = light, Lithoautotroph = medium); every other stream keys
# on substrate alone. This is the single source of truth for stream colour across all figures.
CO2_LIGHT = "#9CC9E8"
CO2_MEDIUM = "#4A97CC"
SUBSTRATE_COLORS = {
    "CO2": CO2_MEDIUM,
    "CO": "#1C6091",
    "CH4": "#3A2E80",
    "VFAs": "#1B9E77",
    "Wet_Organics": "#7ECCA0",
    "Cellulose": "#EAC21C",
    "Polyesters": "#E67E22",
    "Polyamides": "#B23B2E",
    "Other": "#7B2D1A",
}
_ROUTE_COLOR_OVERRIDES = {
    ("CO2", "Photoautotroph (aerobic)"): CO2_LIGHT,
    ("CO2", "Lithoautotroph (aerobic)"): CO2_MEDIUM,
}

# --- node registry -------------------------------------------------------------------------------
# Every node name that appears in numbers.xlsx, mapped to its colour. This is the single place the
# names are recorded: the figure scripts draw from node_colors() rather than each keeping its own
# dict, and check_node_names() turns a rename in the workbook into an error instead of a silent
# fallback to a default colour.
TERMINAL_SOLIDS = "#6E6E6E"          # incinerated solid waste; a fate, deliberately outside the stream palette
PROCESS_NODE = "#9E9E9E"             # unit operations that transform rather than carry carbon
NONCARBON_GREY = "#D0D0D0"           # the non-carbon remainder of an uplinked category

# Logistics input categories: an olive ramp graded by lightness, so the inputs read as off-palette
# against the vivid stream colours. Edit here to restyle every figure that draws the categories.
LOGISTICS_SLATE = {
    "Food": "#9FB200",
    "Clothing": "#758C00",
    "Personal Supplies": "#4E6500",
    "Packaging": "#2A3800",
}
# Metabolic intermediates that are not themselves substrate pools.
METABOLIC_NODES = {
    "Food Waste": "#AA4499",
    "Human Waste": "#857AB8",
}


def node_colors() -> dict[str, str]:
    """Colour for every node name used in numbers.xlsx: Sankey nodes, plus the doughnut's carbon and
    non-carbon split of each logistics category. Stream nodes resolve through SUBSTRATE_COLORS, so a
    change to the substrate palette propagates to every figure."""
    nodes = {
        "Exhaled CO2": SUBSTRATE_COLORS["CO2"],
        "Scrubbed CO2": SUBSTRATE_COLORS["CO2"],
        "Waste CO2 (Vented)": SUBSTRATE_COLORS["CO2"],
        "Sabatier": SUBSTRATE_COLORS["CO2"],
        "Waste Methane (Vented)": SUBSTRATE_COLORS["CH4"],
        "Wet Organics": SUBSTRATE_COLORS["Wet_Organics"],
        "Cellulose": SUBSTRATE_COLORS["Cellulose"],
        "Biodegradable Plastic": SUBSTRATE_COLORS["Polyesters"],
        "Non-Biodegradable Plastic": SUBSTRATE_COLORS["Other"],
        "Solid Carbon Waste": TERMINAL_SOLIDS,
    }
    nodes.update(METABOLIC_NODES)
    for category, base in LOGISTICS_SLATE.items():
        nodes[category] = base
        nodes["Carbon " + category] = base
        nodes["Non-Carbon " + category] = NONCARBON_GREY
    return nodes


def sankey_node_names() -> list[str]:
    """Every distinct node name in the Figure 1A flow table, for validation against the registry."""
    import openpyxl
    ws = openpyxl.load_workbook(NUMBERS, data_only=True, read_only=True)["Figure 1A"]
    names = set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        for cell in row[:2]:
            if cell is not None:
                names.add(str(cell))
    return sorted(names)


def check_node_names(names, context: str = "numbers.xlsx") -> None:
    """Fail if a node name is not in the registry. Without this a renamed node silently takes a
    default colour, which is hard to spot in a rendered figure."""
    unknown = sorted({str(n) for n in names if n is not None} - set(node_colors()))
    if unknown:
        raise KeyError(f"{context}: node name(s) not in common.node_colors(): {unknown}. "
                       f"Add them to the registry or correct the workbook.")



def route_color(substrate: str, mechanism: str, default: str = "#96A1B1") -> str:
    """Fill colour for one route: CO2 splits by mechanism into two shades; every other stream keys on substrate."""
    return _ROUTE_COLOR_OVERRIDES.get((str(substrate), str(mechanism).strip()), SUBSTRATE_COLORS.get(str(substrate), default))


def substrate_color(substrate: str, default: str = "#96A1B1") -> str:
    """Substrate-level colour (CO2 = its representative medium shade); for legends and aggregate views."""
    return SUBSTRATE_COLORS.get(str(substrate), default)


def substrate_legend_items() -> list[tuple[str, str]]:
    return [
        ("CO2", "CO$_2$"),
        ("CO", "CO"),
        ("CH4", "CH$_4$"),
        ("VFAs", "VFAs"),
        ("Wet_Organics", "Wet organics"),
        ("Cellulose", "Cellulose"),
        ("Polyesters", "Polyesters"),
        ("Polyamides", "Polyamides"),
    ]

# Routes that interconvert a stream rather than assimilate it: their carbon is accounted at the route that finally fixes it, so the workbook nulls their throughput, but they still carry a reactor volume. Listed explicitly rather than inferred from a null throughput, which is indistinguishable from a scenario zero.
INTERCONVERSION_ROUTES = frozenset({("Wet_Organics", "Heterotroph (anaerobic)")})


def is_interconversion(row: dict[str, float | str]) -> bool:
    """True if the route interconverts its substrate rather than assimilating carbon from it."""
    return (str(row["substrate"]), str(row["mechanism"])) in INTERCONVERSION_ROUTES


def process_legend_items() -> list[tuple[str, str]]:
    """Organism/process classes that can appear on a hatched (carbon-proportional) area.

    Interconversions are absent: their carbon is accounted at the route that finally fixes it, so they have
    no hatched area and a key entry would advertise a pattern that never renders.
    """
    return [
        ("Phototroph", "\\\\\\"),
        ("Lithotroph", "|||"),
        ("Heterotroph", "///"),
        ("Mixotroph", "xxx"),
        ("Methanotroph", "---"),
        ("Enzymatic", "..."),
        ("Biochemical", "ooo"),
    ]


# --------------------------------------------------------------------------- #
# Manifest
# --------------------------------------------------------------------------- #

def write_manifest(sources: list[Path]) -> None:
    """SHA-256 checksums of every input and output present. Written by build_all.py; a single-figure run leaves any existing manifest untouched."""

    def record(path: Path) -> dict[str, object]:
        return {"path": rel(path), "bytes": path.stat().st_size, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}

    present = [path for path in sources if path.exists()]
    missing = [rel(path) for path in sources if not path.exists()]
    outputs = sorted(OUTPUTS.glob("Fig[0-5].pdf")) + sorted(OUTPUTS.glob("Fig[0-5].png"))
    manifest = {"source_files": [record(p) for p in present], "missing_sources": missing, "outputs": [record(p) for p in outputs]}
    (OUTPUTS / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")